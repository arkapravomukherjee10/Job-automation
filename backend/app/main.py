import io
import json
import os
from datetime import datetime, timezone
from typing import Literal
from uuid import uuid4

import jwt
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Header
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field, HttpUrl, model_validator, field_validator
from sqlalchemy import create_engine, text
from sqlalchemy.pool import StaticPool
from .validation import Resume, validate_resume
from .parsing import prefill_resume

MODE = os.getenv("APP_MODE", "production")
DB = os.getenv("DATABASE_URL", "sqlite:///./jobradar.db")
if MODE != "development" and (
    DB.startswith("sqlite")
    or not os.getenv("SUPABASE_URL")
    or not os.getenv("INVITED_EMAILS")
):
    raise RuntimeError(
        "Production requires PostgreSQL, SUPABASE_URL and INVITED_EMAILS."
    )
engine = create_engine(
    DB,
    **(
        {
            "connect_args": {"check_same_thread": False},
            **({"poolclass": StaticPool} if ":memory:" in DB else {}),
        }
        if DB.startswith("sqlite")
        else {"pool_pre_ping": True}
    ),
)
with engine.begin() as c:
    c.execute(
        text(
            "CREATE TABLE IF NOT EXISTS records (id VARCHAR(36) PRIMARY KEY, owner VARCHAR(100) NOT NULL, kind VARCHAR(30) NOT NULL, data TEXT NOT NULL, created VARCHAR(40) NOT NULL)"
        )
    )
    c.execute(
        text("CREATE INDEX IF NOT EXISTS records_owner_kind ON records(owner, kind)")
    )
    if not DB.startswith("sqlite"):
        c.execute(text("ALTER TABLE records ENABLE ROW LEVEL SECURITY"))
        c.execute(text("REVOKE ALL ON records FROM anon, authenticated"))
app = FastAPI(title="Job Radar API", version="0.1.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=os.getenv("FRONTEND_ORIGINS", "http://localhost:5173").split(","),
    allow_methods=["GET", "POST", "PATCH", "DELETE"],
    allow_headers=["Authorization", "Content-Type"],
)
issuer = os.getenv("SUPABASE_URL", "").rstrip("/") + "/auth/v1"
jwks = (
    jwt.PyJWKClient(issuer + "/.well-known/jwks.json", cache_keys=True)
    if MODE != "development"
    else None
)


def user(authorization: str | None = Header(default=None)):
    if MODE == "development":
        return "local-demo"
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(401, "Sign in required")
    token = authorization[7:]
    try:
        key = jwks.get_signing_key_from_jwt(token)
        claims = jwt.decode(
            token,
            key.key,
            algorithms=["ES256", "RS256"],
            audience="authenticated",
            issuer=issuer,
            options={"require": ["exp", "sub", "email"]},
        )
    except Exception:
        raise HTTPException(401, "Invalid or expired session")
    allowed = {
        x.strip().casefold()
        for x in os.getenv("INVITED_EMAILS", "").split(",")
        if x.strip()
    }
    if claims["email"].casefold() not in allowed:
        raise HTTPException(403, "This beta is invitation only")
    return claims["sub"]


def records(owner, kind):
    with engine.connect() as c:
        rows = (
            c.execute(
                text(
                    "SELECT id, data, created FROM records WHERE owner=:owner AND kind=:kind ORDER BY created DESC"
                ),
                {"owner": owner, "kind": kind},
            )
            .mappings()
            .all()
        )
    return [
        {**json.loads(r["data"]), "id": r["id"], "created": r["created"]} for r in rows
    ]


def add(owner, kind, data):
    record = {
        "id": str(uuid4()),
        "owner": owner,
        "kind": kind,
        "data": json.dumps(data),
        "created": datetime.now(timezone.utc).isoformat(),
    }
    with engine.begin() as c:
        c.execute(
            text(
                "INSERT INTO records (id,owner,kind,data,created) VALUES (:id,:owner,:kind,:data,:created)"
            ),
            record,
        )
    return {**data, "id": record["id"], "created": record["created"]}


def update(owner, kind, rid, data):
    with engine.begin() as c:
        result = c.execute(
            text(
                "UPDATE records SET data=:data WHERE id=:id AND owner=:owner AND kind=:kind"
            ),
            {"id": rid, "owner": owner, "kind": kind, "data": json.dumps(data)},
        )
        if not result.rowcount:
            raise HTTPException(404, "Record not found")


def audit(owner, event):
    add(owner, "audit", {"event": event})


class Preferences(BaseModel):
    roles: list[str] = Field(default_factory=list, max_length=20)
    locations: list[str] = Field(default_factory=list, max_length=20)
    skills: list[str] = Field(default_factory=list, max_length=100)
    blocked_companies: list[str] = Field(default_factory=list, max_length=100)
    threshold: int = Field(default=90, ge=80, le=100)
    resume_consent: bool = False

    @model_validator(mode="before")
    @classmethod
    def migrate_single_preferences(cls, values):
        values = dict(values)
        for old, new in [("role", "roles"), ("location", "locations")]:
            if new not in values:
                value = values.get(old, "")
                values[new] = [value] if value else []
        return values

    @field_validator("roles", "locations")
    @classmethod
    def clean_options(cls, values):
        result = []
        seen = set()
        for value in values:
            value = value.strip()
            if len(value) > 200:
                raise ValueError("Each preference must be 200 characters or fewer")
            if value and value.casefold() not in seen:
                result.append(value)
                seen.add(value.casefold())
        return result


class JobInput(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    company: str = Field(min_length=1, max_length=200)
    location: str = Field(default="", max_length=200)
    description: str = Field(min_length=10, max_length=40000)
    url: HttpUrl
    salary: str = Field(default="", max_length=100)


class Stage(BaseModel):
    stage: Literal[
        "Ready to Apply",
        "Awaiting Approval",
        "Applied (Manual)",
        "In Review",
        "Interview",
        "Offer",
        "Rejected",
        "Archived",
    ]


class ValidationInput(BaseModel):
    base_id: str
    draft: Resume


def prefs(owner):
    found = records(owner, "preferences")
    return Preferences(**found[0]) if found else Preferences()


def score(job, p):
    haystack = (job["title"] + " " + job["description"]).casefold()
    matched = [s for s in p.skills if s.casefold() in haystack]
    checks = []
    if p.roles and not any(
        role.casefold() in job["title"].casefold() for role in p.roles
    ):
        checks.append("Target role differs")
    if p.locations and not any(
        location.casefold() in job["location"].casefold() for location in p.locations
    ):
        checks.append("Preferred location differs")
    value = round(100 * len(matched) / len(p.skills)) if p.skills else None
    blocked = job["company"].casefold() in {s.casefold() for s in p.blocked_companies}
    stage = (
        "Archived"
        if blocked or len(checks) > 1
        else "Near Misses" if checks else "Ready to Apply"
    )
    if not blocked and not checks and value is not None:
        stage = (
            "Archived"
            if value < 60
            else "Awaiting Approval" if value < p.threshold else "Ready to Apply"
        )
    reason = (
        "Company is blocked"
        if blocked
        else (
            "; ".join(checks)
            if checks
            else (
                f"{len(matched)} of {len(p.skills)} preferred skills mentioned"
                if p.skills
                else "Add your skills in Settings to calculate a score"
            )
        )
    )
    return {
        "score": value,
        "matched_skills": matched,
        "stage": stage,
        "reason": reason,
        "blocked": blocked,
        "filtered": len(checks) > 1,
    }


@app.get("/health")
def health():
    with engine.connect() as c:
        c.execute(text("SELECT 1"))
    return {"status": "ok", "mode": MODE}


@app.get("/api/state")
def state(owner=Depends(user)):
    return {
        "preferences": prefs(owner).model_dump(),
        "jobs": [j for j in records(owner, "job") if not j.get("blocked")],
        "resumes": records(owner, "resume"),
        "activity": records(owner, "audit")[:50],
        "capabilities": {"gmail": False, "ai": False},
    }


@app.post("/api/preferences")
def save_preferences(p: Preferences, owner=Depends(user)):
    previous = records(owner, "preferences")
    if previous:
        update(owner, "preferences", previous[0]["id"], p.model_dump())
    else:
        add(owner, "preferences", p.model_dump())
    for job in records(owner, "job"):
        if not job.get("manual_stage") and job["stage"] in [
            "Ready to Apply",
            "Awaiting Approval",
            "Near Misses",
            "Archived",
        ]:
            job.update(score(job, p))
            update(owner, "job", job["id"], job)
    audit(owner, "Preferences updated; open jobs rescored")
    return p


@app.post("/api/jobs", status_code=201)
def create_job(job: JobInput, owner=Depends(user)):
    data = job.model_dump(mode="json")
    if any(j["url"] == data["url"] for j in records(owner, "job")):
        raise HTTPException(409, "This job link is already in your tracker")
    data.update(score(data, prefs(owner)))
    data["source"] = "Added by you"
    result = add(owner, "job", data)
    audit(owner, "Job added: " + data["company"] + " · " + data["title"])
    return result


@app.patch("/api/jobs/{rid}")
def move_job(rid: str, change: Stage, owner=Depends(user)):
    found = next((j for j in records(owner, "job") if j["id"] == rid), None)
    if not found:
        raise HTTPException(404, "Job not found")
    if change.stage == "Applied (Manual)":
        if found.get("applied_at") or found["stage"] in [
            "Applied (Manual)",
            "In Review",
            "Interview",
            "Offer",
            "Rejected",
        ]:
            raise HTTPException(409, "This job is already recorded as applied")
        found["applied_at"] = datetime.now(timezone.utc).isoformat()
    found["stage"] = change.stage
    found["manual_stage"] = True
    update(owner, "job", rid, found)
    audit(owner, found["company"] + " moved to " + change.stage)
    return found


@app.post("/api/resumes/extract")
async def extract(file: UploadFile = File(...), owner=Depends(user)):
    if not prefs(owner).resume_consent:
        raise HTTPException(403, "Enable resume processing consent first")
    content = await file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "Resume must be 5 MB or smaller")
    name = (file.filename or "").lower()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            if len(reader.pages) > 30:
                raise ValueError("Too many pages")
            extracted = "\n".join(p.extract_text() or "" for p in reader.pages)
        elif name.endswith(".docx"):
            import zipfile

            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist()) > 20 * 1024 * 1024:
                    raise ValueError("Expanded document too large")
            from docx import Document

            doc = Document(io.BytesIO(content))
            extracted = "\n".join(
                [p.text for p in doc.paragraphs]
                + [c.text for t in doc.tables for r in t.rows for c in r.cells]
            )
        else:
            raise HTTPException(415, "Upload a PDF or DOCX")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            422, "Could not read this document. Try a text-based PDF or DOCX."
        )
    if not extracted.strip():
        raise HTTPException(
            422,
            "No readable text found. Scanned PDFs need OCR, which is not available yet.",
        )
    draft, warnings = prefill_resume(extracted[:50000])
    return {
        "draft": draft,
        "warnings": warnings,
        "text": extracted[:50000],
        "notice": "Your resume details are prefilled. Review and correct them before confirming your base resume.",
    }


@app.post("/api/resumes", status_code=201)
def save_resume(resume: Resume, owner=Depends(user)):
    if not prefs(owner).resume_consent:
        raise HTTPException(403, "Enable resume processing consent first")
    if not resume.bullets or len({b.id for b in resume.bullets}) != len(resume.bullets):
        raise HTTPException(422, "Provide at least one bullet with unique IDs")
    result = add(
        owner,
        "resume",
        {**resume.model_dump(), "version": len(records(owner, "resume")) + 1},
    )
    audit(owner, "Confirmed base resume version " + str(result["version"]))
    return result


@app.post("/api/resumes/validate")
def validate(payload: ValidationInput, owner=Depends(user)):
    if not prefs(owner).resume_consent:
        raise HTTPException(403, "Enable resume processing consent first")
    base = next(
        (r for r in records(owner, "resume") if r["id"] == payload.base_id), None
    )
    if not base:
        raise HTTPException(404, "Base resume not found")
    result = validate_resume(
        Resume(**{k: base[k] for k in ["facts", "skills", "bullets"]}), payload.draft
    )
    audit(
        owner,
        "Resume consistency check " + ("passed" if result["passed"] else "failed"),
    )
    return result


@app.delete("/api/data", status_code=204)
def delete_data(owner=Depends(user)):
    with engine.begin() as c:
        c.execute(text("DELETE FROM records WHERE owner=:owner"), {"owner": owner})
